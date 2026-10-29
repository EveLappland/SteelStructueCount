"""Final Excel export contracts, without any model calls or reference answers."""
from collections import Counter
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from app.entity_report import component_rows, export, write_single_member_sheet


class EntityReportTests(unittest.TestCase):
    def test_excel_single_member_sheet_matches_components_and_preserves_unknown(self):
        entities = [
            dict(id='beam', spec='HEA220', profile='HEA220', length_mm=1234.56,
                 member_cnt=2, component_multiplier=1),
            dict(id='angles', spec='2LF-L65*6', profile='L65*6', length_mm=2345.67,
                 member_cnt=3, component_multiplier=2),
            dict(id='unknown', spec='HEA280', profile='HEA280', length_mm=None,
                 member_cnt=1, component_multiplier=1),
        ]
        for entity in entities:
            entity.update(view_tags=['LOCAL'], status='MODEL_DERIVED', method='INDEPENDENT',
                          observation_ids=[], length_source_id=None, measurement=None, reason='Local fixture')
        report = dict(entities=entities, observation_entity_map=[], views=[], cases=[],
                      summary={}, conservation={})
        with TemporaryDirectory() as folder:
            directory = Path(folder)
            export(directory, report, {})
            with (directory/'stage_three.xlsx').open('rb') as stream:
                book = load_workbook(stream)
                self.addCleanup(book.close)
                sheet = book['单根材料清单']
                self.assertEqual(list(sheet.values)[0], ('PROFILE', 'MATERIAL', 'LENGTH(mm)'))
                rows = list(sheet.values)[1:]
                self.assertEqual(sheet.max_column, 3)
                self.assertEqual(len(rows), sum(e['member_cnt']*e['component_multiplier'] for e in entities))
                self.assertEqual(Counter(rows), Counter({
                    ('HEA220', None, 1234.56): 2,
                    ('L65*6', None, 2345.67): 6,
                    ('HEA280', None, None): 1,
                }))
                self.assertEqual(sheet['C2'].number_format, '0.00')
                self.assertEqual(sheet.freeze_panes, 'A2')
                self.assertEqual(sheet.auto_filter.ref, 'A1:C10')
                self.assertIn('实体台账', book.sheetnames)
                self.assertIn('Quantity Length Summary', book.sheetnames)
                self.assertIn('展开单根台账', book.sheetnames)
            saved = json.loads((directory/'components.json').read_text(encoding='utf-8'))
            self.assertEqual(saved, component_rows(entities))

    def test_refresh_replaces_only_requested_sheet_without_double_expansion(self):
        book = Workbook()
        self.addCleanup(book.close)
        book.active.title = 'existing'
        book.active['A1'] = 'preserve me'
        rows = [dict(profile='L65*6', length_mm=2345.67)]*2
        write_single_member_sheet(book, rows)
        write_single_member_sheet(book, rows[:1])
        self.assertEqual(book.sheetnames, ['existing', '单根材料清单'])
        self.assertEqual(book['existing']['A1'].value, 'preserve me')
        self.assertEqual(list(book['单根材料清单'].values), [
            ('PROFILE', 'MATERIAL', 'LENGTH(mm)'), ('L65*6', None, 2345.67)])


if __name__ == '__main__':
    unittest.main()
