import io,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import create_app
from test_group_dedup import fixture,row
from test_group_reconciliation import resolution


class AutomaticWorkflowTests(unittest.TestCase):
    def run_job(self,*,manual=False,proposal=None):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        app=create_app(Path(self.tmp.name),start_workers=False)
        client=TestClient(app);client.__enter__();self.addCleanup(client.__exit__,None,None,None)
        identity=client.post('/api/jobs',files=[('files',('drawing.pdf',b'%PDF mock','application/pdf'))],
            data={'auto_run':str(not manual).lower(),'api_key':'test-key','model':'test','base_url':'https://example.test'}).json()['id']
        two=fixture([row('p1','p'),row('a1','a')])
        two.update(workflow='image_first',recognition_mode='view_ledger',candidates=[],warnings=[])
        one={'workflow':'image_first','views':two['views'],'warnings':[]}
        proposal=proposal or {'resolutions':[resolution()], 'exclusions':[],'notes':[]}
        with patch('app.image_workflow.prepare_catalog',return_value=one),patch('app.image_workflow.recognize_views',return_value=two),patch('app.pipeline.load_native',return_value=[]),patch('app.pipeline.OpenAI'),patch('app.pipeline.model_call',return_value=proposal) as calls:
            app.state.worker(identity,1)
        return app,client,identity,calls

    def test_default_pipeline_reaches_real_excel_without_any_manual_approval_or_groups(self):
        app,client,identity,calls=self.run_job()
        job=client.get(f'/api/jobs/{identity}').json()
        self.assertTrue(job['auto_run']);self.assertEqual(job['state'],'completed')
        self.assertEqual(job['total'],3);self.assertEqual(calls.call_count,1)
        self.assertEqual(calls.call_args.kwargs['reasoning_effort'],'max')
        three=app.state.store.stage(identity,3)['data']
        self.assertTrue(three['review']['ai_proposal']['applied'])
        self.assertEqual(three['totals']['allocated_observation_units'],6)
        self.assertEqual(three['totals']['pending_observation_units'],0)
        book=load_workbook(io.BytesIO(client.get(job['download_url']).content))
        self.assertEqual(book['Quantity Length Summary']['C2'].value,3)
        self.assertEqual(book['Member Register'].max_row,4)
        self.assertIn('Pending Register',book.sheetnames)
        self.assertFalse(any(h['action']=='edited' for h in app.state.store.snapshots(identity)))
        # An old completed/empty task can be rerun through AI rather than export-only.
        retried=client.post(f'/api/jobs/{identity}/stages/3/retry',json={'revision':1,'auto_run':True})
        self.assertEqual(retried.status_code,200)
        self.assertEqual(app.state.store.stage(identity,3)['status'],'queued')
        self.assertEqual(app.state.store.stage(identity,3)['revision'],2)

    def test_manual_pause_remains_optional(self):
        app,client,identity,calls=self.run_job(manual=True)
        self.assertEqual(client.get(f'/api/jobs/{identity}').json()['state'],'awaiting_review')
        self.assertFalse(calls.called)
        self.assertEqual(len(app.state.store.get(identity)['stages']),1)

    def test_empty_model_result_is_repaired_twice_then_explicit_failure_not_empty_export(self):
        app,client,identity,calls=self.run_job(proposal={'resolutions':[],'exclusions':[],'notes':[]})
        job=client.get(f'/api/jobs/{identity}').json()
        self.assertEqual(job['state'],'failed');self.assertEqual(job['active_stage'],3)
        self.assertEqual(calls.call_count,3)
        self.assertIsNone(job['download_url'])
        self.assertIn('仍有具备长度',job['error_detail'])

