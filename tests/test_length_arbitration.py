"""Pure contracts and saved real replies; no mocked LLM/API conversations."""
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from app.ai_concurrency import view_concurrency
from app.grid_workflow import Task, ledger_rows
from app.length_arbitration import (ArbitrationResponse, EvidencePlan,
                                    LengthArbitration, candidate, fallback_candidate)
from app.length_expert import LengthExpertResponse, assess_reply
from app.length_context import work_package

ROOT = Path(__file__).resolve().parents[1]


class ArbitrationContractTests(unittest.TestCase):
    def test_env_limit_and_explicit_override_and_five_default(self):
        with patch.dict('os.environ', {}, clear=True):
            self.assertEqual(view_concurrency(), 5)
        with patch.dict('os.environ', {'LLM_MAX_CONCURRENCY': '7'}):
            self.assertEqual(view_concurrency(), 7)
            self.assertEqual(view_concurrency({'max_concurrency': 2}), 2)

    def test_roles_use_shared_task_contract_and_unique_request_ids(self):
        for kind, contract in [('length_evidence', EvidencePlan),
                               ('length_verifier', LengthExpertResponse),
                               ('length_arbiter', ArbitrationResponse)]:
            task = Task('drawing-q', kind, {}, 'prompt', round=1)
            self.assertIs(task.contract, contract)
            self.assertIn(kind, task.name)

    def test_candidate_arithmetic_rejects_unknown_and_keeps_combination_identity(self):
        result = candidate('row-q', 'master', {'method': 'DIAGONAL',
            'horizontal_mm': 3, 'vertical_mm': 4, 'evidence': 'pure arithmetic fixture'})
        self.assertEqual(result['length_mm'], 5)
        self.assertTrue(result['selectable'])
        self.assertFalse(candidate('row-q', 'expert', None)['selectable'])
        self.assertFalse(candidate('row-q', 'expert', result['length_input'], errors=['bad reference'])['selectable'])

    def test_fallback_is_deterministic_and_does_not_grant_expert_authority(self):
        first = candidate('row-q', 'master', {'method': 'DIRECT', 'direct_mm': 3, 'evidence': 'local fixture'})
        second = candidate('row-q', 'expert', {'method': 'DIRECT', 'direct_mm': 4, 'evidence': 'local fixture'})
        self.assertEqual(fallback_candidate([first, second]), first)
        self.assertIsNone(fallback_candidate([candidate('row-q', 'expert', None)]))

    def test_known_direct_rows_are_reviewed_when_scope_is_all_without_revealing_old_numbers(self):
        rows = [{'id': 'beam-q', 'spec': 'TEST-PROFILE', 'view_tags': ['PROJECT-Q/DECK'],
                 'length_mm': 123456, 'length_input': {'method': 'DIRECT', 'direct_mm': 123456},
                 'endpoints': ['OLD_ENDPOINT'], 'location': 'OLD_LOCATION'}]
        item = {'rows': rows, 'length_target_ids': ['beam-q'], 'directory': Path('.'),
                'plan': {'width_px': 10, 'height_px': 10, 'tiles': []},
                'context': {}, 'entry': {'image_path': 'source.png'}}
        _, prompt = work_package(item)
        data = json.loads(prompt.split('输入工作包：\n', 1)[1])
        self.assertEqual([m['member_id'] for m in data['target_members']], ['beam-q'])
        self.assertNotIn('123456', prompt)
        self.assertNotIn('OLD_ENDPOINT', prompt)
        self.assertNotIn('OLD_LOCATION', prompt)


class RecordedArbitrationTests(unittest.TestCase):
    def setUp(self):
        root = ROOT / 'tmp/stage12_length_expert_04_test'
        if not (root / 'report.json').exists():
            self.skipTest('requires previously saved real API run')
        self.one = json.loads((root / 'stage_one.json').read_text(encoding='utf-8'))
        data = json.loads((root / 'stage_two.json').read_text(encoding='utf-8'))
        run = json.loads((root / 'report.json').read_text(encoding='utf-8'))
        self.view = self.one['views'][0]
        analysis = data['view_analysis'][0]
        path = ROOT / 'tmp/length_expert_feature_server/data/jobs' / run['job_id'] / analysis['trace_file']
        master = json.loads((path / 'master_result.json').read_text(encoding='utf-8'))
        self.rows = ledger_rows(master, self.view, self.one, [])
        expert = json.loads((path / 'length_expert/result.json').read_text(encoding='utf-8'))
        self.plan = json.loads((path / 'input.json').read_text(encoding='utf-8'))['grid']
        self.assessment = assess_reply(expert, self.rows, [self.plan['width_px'], self.plan['height_px']])

    def test_expert_remains_candidate_and_blind_verifier_has_no_prior_answers(self):
        with tempfile.TemporaryDirectory() as folder:
            item = {'view': self.view, 'rows': deepcopy(self.rows), 'plan': self.plan,
                    'directory': Path(folder), 'entry': {'image_path': 'original.png'},
                    'context': {'current_view_tag': self.view['plane_name'], 'axis_names': []},
                    'status': {}, 'calls': 0}
            from app.audit import AuditLog
            original = deepcopy(item['rows'])
            controller = LengthArbitration(item, [item], self.assessment, Path(folder),
                AuditLog(Path(folder), stage=2), lambda _: None)
            self.assertEqual(item['rows'], original)
            _, _, prompt = controller.package('length_verifier')
            payload = json.loads(prompt.split('输入工作包：\n', 1)[1])
            self.assertNotIn('candidates', payload)
            for row in payload['target_members']:
                self.assertNotIn('length_mm', row)
                self.assertNotIn('endpoints', row)
            _, _, judged = controller.package('length_arbiter')
            self.assertIn('candidates', json.loads(judged.split('输入工作包：\n', 1)[1]))
            # No accidental row deletion even when no judge answer is available.
            controller.failed('length_arbiter', 'deliberate local failure-path check')
            self.assertEqual([(r['id'], r['spec'], r['view_tags'], r['member_cnt']) for r in original],
                             [(r['id'], r['spec'], r['view_tags'], r['member_cnt']) for r in item['rows']])
            self.assertTrue(all(not r['length_arbitration']['manual_required']
                                for r in item['rows'] if r.get('length_arbitration')))

    def test_subset_validation_does_not_demand_other_unknown_rows(self):
        identity = next(iter(self.assessment['updates']))
        update = self.assessment['updates'][identity]
        payload = {'references': self.assessment['payload']['references'],
                   'results': [update['resolution']], 'context_request': None}
        checked = assess_reply(payload, self.rows, [self.plan['width_px'], self.plan['height_px']], target_ids=[identity])
        self.assertEqual(checked['missing'], [])
        self.assertFalse(any(f['check'] == 'target_coverage' for f in checked['feedback']))


class RecordedLiveArbitrationTests(unittest.TestCase):
    def test_export_keeps_final_arbitration_audit_without_advancing_a_job(self):
        path = ROOT / 'tmp/arbitration_02_test/stage_two.json'
        if not path.exists():
            self.skipTest('requires completed real arbitration')
        from app.group_reconciliation import reconcile_exact
        from app.pipeline import export
        from openpyxl import load_workbook
        data = json.loads(path.read_text(encoding='utf-8'))
        count = sum(bool(r.get('length_arbitration')) for r in data['view_members'])
        with tempfile.TemporaryDirectory() as folder:
            export(Path(folder), data, data, reconcile_exact(data), [], [])
            book = load_workbook(Path(folder) / 'output/steel_takeoff.xlsx', read_only=True)
            try:
                sheet = book['Length Arbitration']
                self.assertEqual(sheet.max_row, count + 1)
                self.assertTrue(all('candidates' in str(row[-1]) for row in list(sheet.values)[1:]))
            finally:
                book.close()

    def test_saved_real_judgments_pass_current_guards_and_preserve_every_source_identity(self):
        folder = ROOT / 'tmp/arbitration_02_test'
        if not (folder / 'report.json').exists():
            self.skipTest('requires completed live API arbitration, not a fake response')
        run = json.loads((folder / 'report.json').read_text(encoding='utf-8'))
        one = json.loads((folder / 'stage_one.json').read_text(encoding='utf-8'))
        two = json.loads((folder / 'stage_two.json').read_text(encoding='utf-8'))
        root = ROOT / 'tmp/arbitration_feature_server/data/jobs' / run['job_id']
        from app.member_recognition import key_record, validate_ledger
        validated = validate_ledger(deepcopy(two))
        for row in validated['view_members']:
            self.assertEqual(row['member_key'], key_record(row))
            self.assertIn(row['source_plane_tag'], row['view_tags'])
        for analysis in two['view_analysis']:
            trace = root / analysis['trace_file']
            view = next(v for v in one['views'] if v['id'] == analysis['view_id'])
            rows = ledger_rows(json.loads((trace / 'master_result.json').read_text(encoding='utf-8')), view, one, [])
            final = [r for r in validated['view_members'] if r['source_view'] == view['id']]
            self.assertEqual([(r['id'], r['spec'], r['member_cnt'], r['view_tags']) for r in rows],
                             [(r['id'], r['spec'], r['member_cnt'], r['view_tags']) for r in final])
            saved = json.loads((trace / 'arbitration/result.json').read_text(encoding='utf-8'))
            obj = object.__new__(LengthArbitration)
            obj.targets = list(json.loads((trace / 'arbitration/candidates_initial.json').read_text(encoding='utf-8')))
            obj.candidates = {r['id']: r['length_arbitration']['candidates'] for r in final if r.get('length_arbitration')}
            obj.images = json.loads((trace / 'arbitration/evidence_images.json').read_text(encoding='utf-8'))
            obj.item = {'plan': json.loads((trace / 'input.json').read_text(encoding='utf-8'))['grid']}
            self.assertEqual(obj.check_verdicts(saved['verdict']), [])
            for row in final:
                if row.get('length_arbitration'):
                    self.assertEqual(row['length_arbitration'],
                        next(r for r in validated['view_members'] if r['id'] == row['id'])['length_arbitration'])
