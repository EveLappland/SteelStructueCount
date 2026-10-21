"""Local contracts/reference math only. No fabricated model/HTTP responses."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from PIL import Image
from app.length_expert import calibration,estimate,target_rows,assess_reply,apply_reply,LengthExpertResponse,DimensionReference
from app.length_context import freeze_context,work_package
from app.grid_workflow import Task
from app.member_recognition import validate_ledger


def reference(identity,value,axis='X',start=(.1,.1),end=(.9,.1)):
    return {'id':identity,'kind':'DIMENSION','axis':axis,'value_mm':value,
            'start':{'name':'first','point':start},'end':{'name':'second','point':end},'evidence':'local arithmetic fixture, not a model answer'}


class LengthMathTests(unittest.TestCase):
    def test_independent_calibration_and_conflict(self):
        refs=[reference('D1',8000),reference('D2',4000,start=(.1,.3),end=(.5,.3))]
        scale=calibration(refs,[1000,1000]);self.assertEqual(scale['X']['status'],'consistent')
        self.assertEqual(scale['X']['mm_per_px'],10);self.assertEqual(scale['Y']['status'],'insufficient')
        points=[{'point':(.1,.5)},{'point':(.7,.5)}]
        self.assertEqual(estimate(points,scale,[1000,1000])['length_mm'],6000)
        self.assertIsNone(estimate([{'point':(.1,.1)},{'point':(.7,.7)}],scale,[1000,1000]))
        refs[1]['value_mm']=5000
        self.assertEqual(calibration(refs,[1000,1000])['X']['status'],'conflict')

    def test_same_span_copied_twice_is_not_independent_calibration(self):
        refs=[reference('D1',8000),reference('D2',8000)]
        self.assertEqual(calibration(refs,[1000,1000])['X']['status'],'insufficient')

    def test_no_scale_inferred_from_level_values_or_tiny_spans(self):
        ref=reference('EL1',107500);ref['kind']='LEVEL'
        small=reference('D1',1000,start=(.1,.1),end=(.12,.1))
        self.assertEqual(calibration([ref,small],[1000,1000])['X']['samples'],[])

    def test_level_reference_may_have_only_one_anchor(self):
        ref=DimensionReference.model_validate({'id':'EL1075','kind':'LEVEL','axis':'Y','value_mm':107500,
            'start':{'name':'EL.107.500(TOS)','point':[.1,.5]},'evidence':'single level datum, no span endpoint'})
        self.assertIsNone(ref.end)
        self.assertEqual(calibration([ref.model_dump()],[1000,1000])['Y']['status'],'insufficient')

    def test_context_crop_contains_exact_original_pixels(self):
        with tempfile.TemporaryDirectory() as folder:
            directory=Path(folder);(directory/'views').mkdir();source=directory/'views/source.png'
            original=Image.new('RGB',(200,100),'white');original.putpixel((100,50),(20,40,60));original.save(source,dpi=(300,300))
            crop=freeze_context(directory,'different-project',source,{'bbox':[.1,.1,.9,.9],'reason':'include both ends and remote ruler'})
            with Image.open(directory/'views'/crop['image_file']) as image:self.assertEqual(image.tobytes(),original.crop([20,10,180,90]).tobytes())

    def test_task_uses_same_request_dispatch_contract(self):
        task=Task('v','length_expert',{},'prompt',round=1)
        self.assertIs(task.contract,LengthExpertResponse);self.assertEqual(task.name,'v_length_expert01_format00')

    def test_first_request_hides_old_geometric_conclusions(self):
        rows=[{'id':'m','spec':'HEA260','view_tags':['Q/7'],'length_mm':123456,
               'length_input':{'method':'DIAGONAL','horizontal_mm':765432},
               'endpoints':['WRONG_OLD_ENDPOINT','WRONG_OLD_ENDPOINT2'],'location':'WRONG_OLD_LOCATION',
               'candidate_ids':['label-1']}]
        item={'rows':rows,'plan':{'width_px':1000,'height_px':1000,'tiles':[]},'directory':Path('.'),
              'context':{},'entry':{'image_path':'original.png'},'located':[{'observations':[
                  {'candidate_id':'label-1','label_point_original_px':[300,400],'source_tile':'tile-1'}]}]}
        _,prompt=work_package(item)
        self.assertNotIn('WRONG_OLD_ENDPOINT',prompt);self.assertNotIn('WRONG_OLD_LOCATION',prompt)
        self.assertNotIn('123456',prompt);self.assertNotIn('765432',prompt)
        self.assertIn('point_original_px',prompt);self.assertIn('label-1',prompt)

    def test_targets_unknown_diagonal_and_unresolved_only(self):
        rows=[{'id':'a','length_mm':None},{'id':'b','length_mm':5000,'length_input':{'method':'DIAGONAL'}},
              {'id':'c','length_mm':6000,'segmentation_status':'UNRESOLVED'},{'id':'d','length_mm':8000}]
        self.assertEqual([r['id'] for r in target_rows(rows)],['a','b','c'])

    def test_expression_arithmetic_references_and_no_row_mutation(self):
        # Test the pure local expression engine, not a mocked API conversation.
        rows=[{'id':'m','spec':'HEA260','view_tags':['Q/7'],'source_view':'v','member_cnt':1,'length_mm':None},
              {'id':'keep','spec':'HEA240','view_tags':['Q/8'],'member_cnt':1,'length_mm':8000}]
        expression={'member_id':'m','status':'CONFIRMED','method':'DIAGONAL',
            'endpoints':[{'name':'node A','point':[.1,.1]},{'name':'node B','point':[.5,.5]}],
            'components':[{'direction':'HORIZONTAL','terms':[{'reference_id':'D1','factor':.5}],'basis':'explicit half-span'},
                          {'direction':'VERTICAL','terms':[{'reference_id':'D2','factor':1}],'basis':'dimension'}],
            'reference_ids':['D1','D2'],'issue':None}
        local={'references':[reference('D1',6000),reference('D2',4000,'Y',(.1,.1),(.1,.5))],
               'results':[expression],'context_request':None}
        original=copy.deepcopy(rows);assessment=assess_reply(local,rows,[1000,1000])
        self.assertEqual(assessment['feedback'],[])
        after,changes=apply_reply(rows,assessment)
        self.assertEqual(rows,original);self.assertEqual(after[1],original[1]);self.assertEqual(after[0]['length_mm'],5000)
        self.assertEqual(after[0]['length_input']['horizontal_mm'],3000)
        self.assertEqual(after[0]['member_cnt'],1);self.assertEqual(after[0]['view_tags'],['Q/7'])
        self.assertTrue(changes[0]['applied'])
        local['results'][0]['components'][0]['terms'][0]['reference_id']='invented'
        bad=assess_reply(local,rows,[1000,1000]);rejected,_=apply_reply(rows,bad)
        self.assertTrue(bad['feedback']);self.assertIsNone(rejected[0]['length_mm'])
        local['results'][0]['member_id']='not-a-target'
        self.assertEqual(assess_reply(local,rows,[1000,1000])['missing'],['m'])


class RecordedExpertTests(unittest.TestCase):
    def test_real_05_reply_accepts_single_point_level_references(self):
        root=Path(__file__).resolve().parents[1]
        directory=root/'tmp/length_expert_feature_server/data/jobs/a773986b84a0413ca1fe597912902e0d'
        paths=list((directory/'responses').glob('attempt-grid-*/f000-p001-view03_length_expert00_format00.json'))
        if not paths:self.skipTest('需要保存的05真实专家回复')
        from app.pipeline import response_content
        payload=LengthExpertResponse.model_validate_json(response_content(json.loads(paths[0].read_text(encoding='utf-8'))))
        self.assertEqual(len(payload.results),4)
        self.assertTrue(any(r.kind=='LEVEL' and r.end is None for r in payload.references))

    def test_real_expert_reply_preserves_identity_and_survives_draft_validation(self):
        root=Path(__file__).resolve().parents[1]
        directory=root/'tmp/length_expert_feature_server/data/jobs/4afc8ea11a3a49bb83b83ce64e0a715a'
        paths=list((directory/'grid_runs').glob('*/f000-p001-view02/length_expert/result.json'))
        if not paths:self.skipTest('需要保存的真实长度专家回复')
        before=json.loads((directory/'acceptance/before.json').read_text(encoding='utf-8'))
        rows=[r for r in before['view_members'] if r['source_view']=='f000-p001-view02']
        trace=json.loads((paths[0].parents[1]/'input.json').read_text(encoding='utf-8'))
        plan=trace['grid'];payload=json.loads(paths[0].read_text(encoding='utf-8'))
        assessment=assess_reply(payload,rows,[plan['width_px'],plan['height_px']])
        self.assertFalse(assessment['feedback'])
        after,changes=apply_reply(rows,assessment)
        self.assertEqual([(r['id'],r['spec'],r['view_tags'],r['member_cnt']) for r in rows],
                         [(r['id'],r['spec'],r['view_tags'],r['member_cnt']) for r in after])
        self.assertEqual(len(changes),12)
        combined={**before,'view_members':after}
        validated=validate_ledger(combined)['view_members']
        for old,new in zip(after,validated):
            if old.get('length_review'):
                self.assertEqual(old['length_review'],new['length_review'])
                self.assertEqual(new['length_mm'],5000)
                self.assertEqual(new['length_input']['horizontal_mm'],3000)
                self.assertEqual(new['length_input']['vertical_mm'],4000)

    def test_offline_export_includes_real_expert_history(self):
        root=Path(__file__).resolve().parents[1]
        path=root/'tmp/length_expert_feature_server/data/jobs/4afc8ea11a3a49bb83b83ce64e0a715a/acceptance/after.json'
        if not path.exists():self.skipTest('需要保存的真实专家台账')
        from app.group_reconciliation import reconcile_exact
        from app.pipeline import export
        from openpyxl import load_workbook
        data=json.loads(path.read_text(encoding='utf-8'))
        # Pure offline export in a disposable directory: no job state/API and
        # no third stage is created for the user's experiment.
        with tempfile.TemporaryDirectory() as folder:
            export(Path(folder),data,data,reconcile_exact(data),[],[])
            book=load_workbook(Path(folder)/'output/steel_takeoff.xlsx',read_only=True)
            try:
                sheet=book['Length Expert Audit'];self.assertEqual(sheet.max_row,31)
                values=list(sheet.values)
                self.assertTrue(any(row[4] is None and row[5]==5500 for row in values[1:]))
                self.assertTrue(any('reference_ids' in str(row[6]) for row in values[1:]))
            finally:book.close()
