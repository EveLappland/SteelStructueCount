import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pymupdf as fitz
from fastapi.testclient import TestClient
from openai.types.chat import ChatCompletion
from openpyxl import load_workbook

from app.main import create_app
from app.native_pdf import extract_page
from app.pipeline import (ViewResponse, AnalysisResponse, model_call, parse_model_content, validate_draft,
                          evidence_aliases, repair_model_evidence)
from app.pipeline import model_schema, request_options, BoundaryResponse, apply_boundary_reviews
from app.store import JobStore, StoreError
from app.discovery_inventory import build_inventory
from test_entities import fixture, observation


def pdf_bytes():
    document=fitz.open()
    page=document.new_page(width=600,height=400)
    page.insert_text((30,340),'PLAN AT EL.100.000(TOS)',fontsize=14)
    page.insert_text((310,340),'ELEVATION ON ALIGNMENT A',fontsize=14)
    page.insert_text((40,50),'8000 16000 6000',fontsize=12)
    page.insert_text((40,70),'HEA260',fontsize=12)
    page.draw_line((40,100),(240,100))
    page.draw_line((310,100),(510,100))
    result=document.tobytes()
    document.close()
    return result


def prepare_mock_data():
    data=fixture([observation('a'),observation('b','elevation-A')])
    # Extraction's first text is stable for this deliberately small vector fixture.
    result=json.loads(json.dumps(data).replace('"e"','"f000-p001:t2"').replace('"p"','"f000-p001"'))
    result['frame']['origin_evidence_ids']=['f000-p001:t0']
    for o in result['observations']:
        for p in ('start','end'):o[p]['z']['evidence_ids']=['f000-p001:t0']
    return result


class StoreTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.store=JobStore(Path(self.tmp.name)/'jobs.sqlite3')
        self.store.create('job')
        self.store.claim('job',1)
        self.store.complete_stage('job',1,{'draft':1})

    def tearDown(self):self.tmp.cleanup()

    def test_stale_versions_and_repeated_approval(self):
        self.store.edit('job',1,1,{'draft':2},'correct endpoint')
        with self.assertRaises(StoreError):self.store.approve('job',1,1)
        self.assertEqual(self.store.approve('job',1,2),2)
        self.assertIsNone(self.store.approve('job',1,2))
        self.assertEqual(len(self.store.get('job')['stages']),2)

    def test_upstream_edits_invalidate_downstream_and_hide_export(self):
        self.store.approve('job',1,1)
        self.store.claim('job',2);self.store.complete_stage('job',2,{'length':1000})
        self.store.approve('job',2,1)
        self.store.claim('job',3);self.store.complete_stage('job',3,{'count':1})
        self.store.approve('job',3,1)
        self.store.claim('job',4);self.store.finish('job',{'workbook':'old.xlsx'})
        self.store.edit('job',1,1,{'draft':3},'change length evidence')
        self.assertEqual(self.store.stage('job',2)['status'],'stale')
        self.assertEqual(self.store.stage('job',3)['status'],'stale')
        self.assertNotIn('workbook',self.store.get('job'))
        self.assertTrue(any(s['data']=={'count':1} for s in self.store.snapshots('job')))

    def test_restart_marks_busy_work_interrupted(self):
        self.store.approve('job',1,1);self.store.claim('job',2)
        reopened=JobStore(self.store.path);reopened.recover()
        self.assertEqual(reopened.get('job')['state'],'interrupted')
        self.assertEqual(reopened.stage('job',1)['status'],'approved')
        self.assertEqual(reopened.retry('job',2,1),2)
        self.assertEqual(reopened.stage('job',2)['revision'],2)
        self.assertTrue(any(h['action']=='retry' for h in JobStore(self.store.path).stage('job',2)['history']))

    def test_retry_never_overwrites_manual_draft(self):
        self.store.edit('job',1,1,{'draft':'human'},'manual correction')
        self.store.retry('job',1,2);self.store.claim('job',1)
        self.store.complete_stage('job',1,{'draft':'model retry'})
        self.assertEqual(self.store.stage('job',1)['data'],{'draft':'human'})
        self.assertTrue(any(s['action']=='retry_proposal' for s in self.store.snapshots('job')))

    def test_export_failure_retries_export_only(self):
        for stage in (1,2,3):
            if stage>1:self.store.claim('job',stage);self.store.complete_stage('job',stage,{'draft':stage})
            self.store.approve('job',stage,1)
        self.store.claim('job',4);self.store.fail('job',4,'disk error')
        self.assertEqual(self.store.retry('job',3,1),4)
        self.assertEqual(self.store.stage('job',3)['status'],'approved')


class ApiWorkflowTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
        self.app=create_app(self.root,start_workers=False)
        self.client=TestClient(self.app)
        self.client.__enter__()
        self.job_id=self.client.post('/api/jobs',files=[('files',('drawing.pdf',pdf_bytes(),'application/pdf'))],
            data={'api_key':'secret-test-key','base_url':'https://example.test','model':'test','auto_run':'false'}).json()['id']
        self.base=f'/api/jobs/{self.job_id}'
        data=prepare_mock_data()
        def mock_call(client,model,prompt,entry,response_path,contract,**kwargs):
            response_path.parent.mkdir(parents=True,exist_ok=True)
            response_path.write_text('{"mock":true}',encoding='utf-8')
            if contract is ViewResponse:return {'views':data['views']}
            from app.stage_one import PageAssessment,ConstructionCatalog
            if contract is PageAssessment:return {'decision':'CONSTRUCTION','reason':'包含平面和立面建造布置图。'}
            if contract is ConstructionCatalog:
                return {'views':[{'title':v['title'],'name':v.get('fixed_label') or 'EL.100.000(TOS)',
                    'kind':v['kind'],'bbox':v['bbox'],'plane_names':['A','EL.100.000(TOS)'],
                    'steel_types':['HEA240','HEA260']} for v in data['views']]}
            from app.member_recognition import MemberResponse
            if contract is MemberResponse:
                import json
                bundle=json.loads(prompt.split('输入包：\n',1)[1])
                return {'members':[{'spec':'HEA240','view_tags':[v['tag'] for v in bundle['view_catalog']],
                                    'member_cnt':1,'length_mm':6000,'location':'A轴与B轴之间','issue':None}]}
            view_id='plan' if '-view01' in str(response_path) else 'elevation-A'
            selected=[o for o in data['observations'] if o['view_id']==view_id]
            return {'frame':data['frame'],'observations':selected,'axis_marks':[],'warnings':[]}
        # Legacy drafts remain valid for geometry/version/export regressions.
        self.mock_call=mock_call
        self.legacy_draft=data
        store=self.app.state.store;store.claim(self.job_id,1)
        seed=self._legacy_catalog(self.app.state.jobs_directory/self.job_id,data,lambda _:None,None)
        store.complete_stage(self.job_id,1,seed)

    def _legacy_catalog(self,directory,data,emit,audit):
        from app.pipeline import extract
        extract(directory,emit,audit)
        return validate_draft(directory,1,data)

    def test_image_catalog_then_recognition_and_upstream_invalidation(self):
        from app.member_recognition import MemberResponse
        from app.stage_one import PageAssessment,ConstructionCatalog
        self.legacy_draft['views'][1]['title']='ELEVATION ON ALIGNMENT A'
        for index,view in enumerate(self.legacy_draft['views']):
            view['bbox']=[0,0,.49,1] if index==0 else [.51,0,1,1]
            view['structure_bbox']=[.06,.2,.4,.4] if index==0 else [.55,.2,.9,.4]
        identity=self.client.post('/api/jobs',files=[('files',('drawing.pdf',pdf_bytes(),'application/pdf'))],
                                  data={'base_url':'https://example.test','api_key':'catalog-test-key','model':'test','max_concurrency':'2','auto_run':'false'}).json()['id']
        base=f'/api/jobs/{identity}'
        with patch('app.pipeline.OpenAI'),patch('app.pipeline.model_call',side_effect=self.mock_call) as model:
            self.app.state.worker(identity,1)
            self.assertTrue(model.called)
            self.assertEqual([call.args[5] for call in model.call_args_list],[PageAssessment,ConstructionCatalog])
            self.assertTrue(all(call.kwargs['reasoning_effort']=='high' for call in model.call_args_list))
        stage=self.client.get(base+'/stages/1').json()
        self.assertEqual(stage['status'],'awaiting_review')
        self.assertEqual(stage['data']['workflow'],'image_first')
        self.assertEqual(stage['data']['observations'],[])
        self.assertTrue(all(v['image_file'] for v in stage['data']['views']))
        image=stage['data']['views'][0]['image_file']
        original=self.client.get(base+'/views/'+image).content
        self.assertTrue(original.startswith(b'\x89PNG'))
        self.assertEqual(self.client.get(base+'/stages/2').status_code,404)
        invalid=self.client.post(base+'/stages/1/approve',json={'revision':1,'acknowledge_pending':True,
            'settings':{'max_concurrency':'0'}})
        self.assertEqual(invalid.status_code,422)
        self.assertEqual(self.client.get(base+'/stages/1').json()['status'],'awaiting_review')
        self.assertEqual(self.client.get(base+'/stages/2').status_code,404)
        response=self.client.post(base+'/stages/1/approve',json={'revision':1,'acknowledge_pending':True,
            'settings':{'base_url':'https://example.test','api_key':'recognition-test-key','model':'test'}})
        self.assertEqual(response.status_code,200,response.text)
        with patch('app.pipeline.OpenAI'),patch('app.pipeline.model_call',side_effect=self.mock_call) as model:
            self.app.state.worker(identity,2)
            self.assertTrue(model.called)
            self.assertTrue(all(call.args[5] is MemberResponse for call in model.call_args_list))
            self.assertTrue(all(call.args[3]['image_path'].endswith('_300dpi.png') for call in model.call_args_list))
            self.assertTrue(all('native_evidence' in call.args[2] for call in model.call_args_list))
            self.assertTrue(all('profile_catalog' in call.args[2] for call in model.call_args_list))
        stage_two=self.client.get(base+'/stages/2').json()
        self.assertEqual(stage_two['status'],'awaiting_review')
        batches=[r for r in self.client.get(base+'/logs?limit=1000').json()['items'] if r['event']=='VIEW_BATCH_STARTED']
        self.assertEqual(batches[-1]['configured_concurrency'],2)
        self.assertEqual(len(stage_two['data']['view_members']),2)
        self.assertTrue(all(r['location']=='A轴与B轴之间' and 'basis' not in r for r in stage_two['data']['view_members']))
        self.assertTrue(any(a.get('coverage_differences') for a in stage_two['data']['view_analysis']))
        self.assertEqual(self.client.get(base+'/views/'+image).content,original)
        invalid=self.client.post(base+'/stages/2/retry',json={'revision':stage_two['revision'],
            'settings':{'max_concurrency':'-1'}})
        self.assertEqual(invalid.status_code,422)
        self.assertEqual(self.client.get(base+'/stages/2').json()['revision'],stage_two['revision'])
        retry=self.client.post(base+'/stages/2/retry',json={'revision':stage_two['revision']})
        self.assertEqual(retry.status_code,200,retry.text)
        with patch('app.pipeline.OpenAI'),patch('app.pipeline.model_call',side_effect=self.mock_call) as model:
            self.app.state.worker(identity,2)
            self.assertTrue(model.called)
            self.assertTrue(all(call.kwargs['reasoning_effort']=='max' for call in model.call_args_list))
        stage_two=self.client.get(base+'/stages/2').json()
        edited=stage_two['data'];edited['view_members'][0]['length_mm']=5000
        edit=self.client.patch(base+'/stages/2',json={'revision':stage_two['revision'],'data':edited,'reason':'尺寸人工修正'})
        self.assertEqual(edit.status_code,200,edit.text)
        self.assertEqual(edit.json()['data']['view_members'][0]['length_mm'],5000)
        stale=self.client.patch(base+'/stages/2',json={'revision':stage_two['revision'],'data':edited,'reason':'旧版本'})
        self.assertEqual(stale.status_code,409)
        tampered=json.loads(json.dumps(edit.json()['data']));tampered['views'][0]['title']='UNCONFIRMED PLAN'
        rejected=self.client.patch(base+'/stages/2',json={'revision':edit.json()['revision'],'data':tampered,'reason':'试图篡改目录'})
        self.assertEqual(rejected.status_code,422)
        # The image ledger now reaches a separate group reconciliation stage.
        response=self.client.post(base+'/stages/2/approve',json={'revision':edit.json()['revision'],'acknowledge_pending':True})
        self.assertEqual(response.status_code,200,response.text)
        from app.group_reconciliation import reconcile
        def automatic_proposal(two,*args,**kwargs):
            data=reconcile(two)
            data['resolutions']=[{'id':f'auto-{i}','spec':r['spec'],'view_tags':r['view_tags'],
                'length_mm':r['length_mm'],'member_cnt':r['member_cnt'],
                'allocations':[{'row_id':r['id'],'member_cnt':r['member_cnt']}],
                'reason':'模拟AI：无法确定重复关系，独立保留'} for i,r in enumerate(two['view_members'])]
            data=reconcile(two,data)
            return {'proposal':{'resolutions':data['resolutions'],'exclusions':[],'notes':['待原图核对']},
                    'validation':{'valid':True},'data':data,'folder':'responses/test'}
        with patch('app.group_proposal.propose_groups',side_effect=automatic_proposal):self.app.state.worker(identity,3)
        third=self.client.get(base+'/stages/3').json()
        self.assertEqual(third['data']['reconciliation_mode'],'group_ledger')
        self.assertTrue(third['data']['members'])
        self.assertTrue(third['data']['review']['ai_proposal']['applied'])
        ai_retry=self.client.post(base+'/stages/3/retry',json={'revision':third['revision'],
            'reuse_saved':False,'settings':{'review_mode':'ai'}})
        self.assertEqual(ai_retry.status_code,200,ai_retry.text)
        with patch('app.group_proposal.propose_groups',side_effect=automatic_proposal) as propose:
            self.app.state.worker(identity,3)
            self.assertEqual(propose.call_args.args[2]['review_mode'],'ai')
            self.assertFalse(propose.call_args.kwargs['reuse_saved'])
        third=self.client.get(base+'/stages/3').json()
        self.assertEqual(third['data']['review']['ai_proposal']['proposal']['notes'],['待原图核对'])
        self.assertTrue(third['data']['resolutions'])
        from app.member_recognition import view_tag
        group=third['data'];source=edit.json()['data']
        group['review']['ai_proposal']['proposal']['notes']=['伪造的模型说明']
        group['resolutions']=[{'id':'verified','spec':'HEA260',
            'view_tags':[view_tag(v) for v in source['views']], 'length_mm':6000,'member_cnt':1,
            'allocations':[{'row_id':r['id'],'member_cnt':1} for r in source['view_members']],
            'reason':'测试夹具：纠正模型规格和长度后确认两个投影为同一实体'}]
        saved=self.client.patch(base+'/stages/3',json={'revision':third['revision'],'data':group,'reason':'独立核对夹具'})
        self.assertEqual(saved.status_code,200,saved.text)
        self.assertEqual(saved.json()['data']['summary'][0]['quantity'],1)
        self.assertEqual(saved.json()['data']['review']['ai_proposal']['proposal']['notes'],['待原图核对'])
        # Coverage discrepancies remain visible even after all rows are assigned.
        self.assertTrue(saved.json()['data']['exceptions'])
        approved=self.client.post(base+'/stages/3/approve',json={'revision':saved.json()['revision'],'acknowledge_pending':True})
        self.assertEqual(approved.status_code,200,approved.text)
        self.app.state.worker(identity,4)
        completed=self.client.get(base).json()
        self.assertEqual(completed['state'],'completed')
        self.assertEqual(completed['total'],1)
        book=load_workbook(io.BytesIO(self.client.get(completed['download_url']).content))
        self.assertEqual(book['Quantity Length Summary']['C2'].value,1)
        self.assertIn('Group Allocations',book.sheetnames)
        self.assertIn('Excluded Observations',book.sheetnames)
        again=self.client.post(base+'/stages/3/approve',json={'revision':saved.json()['revision']})
        self.assertEqual(again.status_code,200,again.text)
        self.assertEqual(self.client.get(base).json()['total'],1)
        upstream=stage['data'];upstream['views'][0]['axis_labels']=['PROJECT-Q/9'];upstream['views'][0]['axes_reviewed']=True
        edit=self.client.patch(base+'/stages/1',json={'revision':stage['revision'],'data':upstream,'reason':'人工确认轴线'})
        self.assertEqual(edit.status_code,200,edit.text)
        self.assertEqual(self.client.get(base+'/stages/2').json()['status'],'stale')
        self.assertEqual(self.client.get(base+'/stages/3').json()['status'],'stale')
        self.assertIsNone(self.client.get(base).json().get('download_url'))

    def tearDown(self):
        self.client.__exit__(None,None,None)
        self.tmp.cleanup()

    def stage(self,n):return self.client.get(self.base+f'/stages/{n}').json()

    def test_first_stage_returns_catalog_counts_and_original_evidence(self):
        stage=self.stage(1)
        inventory=stage['data']['discovery_inventory']
        self.assertEqual(inventory['profile_count'],1)
        self.assertEqual(inventory['profiles'][0]['profile'],'HEA260')
        self.assertEqual(inventory['profiles'][0]['evidence'][0]['text'],'HEA260')
        self.assertEqual(inventory['view_counts']['PLAN'],1)
        self.assertEqual(inventory['view_counts']['ELEVATION'],1)
        self.assertEqual(stage['revision'],self.stage(1)['revision'])

    def approve(self,n):
        stage=self.stage(n)
        response=self.client.post(self.base+f'/stages/{n}/approve',json={'revision':stage['revision'],'acknowledge_pending':True})
        self.assertEqual(response.status_code,200,response.text)
        self.app.state.worker(self.job_id,n+1)

    def test_mocked_end_to_end_pauses_and_exports(self):
        self.assertEqual(self.client.get(self.base).json()['state'],'awaiting_review')
        self.assertEqual(self.client.get(self.base+'/stages/2').status_code,404)
        stage=self.stage(1)
        self.assertEqual(len(stage['data']['views']),2)
        self.assertEqual(len(stage['data']['observations']),2)
        self.approve(1)
        self.assertEqual(self.stage(2)['status'],'awaiting_review')
        self.assertTrue(all(c['confirmed'] for c in self.stage(2)['data']['candidates']))
        self.approve(2)
        self.assertEqual(len(self.stage(3)['data']['members']),1)
        self.approve(3)
        status=self.client.get(self.base).json()
        self.assertEqual(status['total'],1)
        self.assertEqual(status['unresolved'],0)
        response=self.client.get(status['download_url'])
        self.assertEqual(response.status_code,200)
        trace=self.client.get(self.base+'/logs').json()
        self.assertTrue(any(r['event']=='LENGTH_EVALUATED' for r in trace['items']))
        self.assertTrue(any(r['event']=='EXPORT_COMPLETED' for r in trace['items']))
        self.assertNotIn('secret-test-key',json.dumps(trace))
        self.assertEqual(self.client.get(self.base+'/logs-download').status_code,200)
        book=load_workbook(io.BytesIO(response.content))
        self.assertEqual(book['Quantity Length Summary']['D2'].value,1)
        self.assertEqual(book['Member Register']['H2'].value,8000)
        book.close()
        archive=self.client.get(status['views_download_url'])
        with zipfile.ZipFile(io.BytesIO(archive.content)) as package:
            self.assertIn('evidence.json',package.namelist())
            self.assertIn('stage_3.json',package.namelist())
            history=package.read('review_history.json').decode()
            self.assertNotIn('secret-test-key',history)
        repeated=self.client.post(self.base+'/stages/3/approve',json={'revision':1})
        self.assertEqual(repeated.status_code,200)
        self.assertEqual(self.client.get(self.base).json()['state'],'completed')

    def test_derived_lengths_cannot_be_forged_and_versions_conflict(self):
        self.approve(1)
        stage=self.stage(2)
        stage['data']['candidates'][0]['length_mm']=99999
        stage['data']['candidates'][0]['confirmed']=True
        response=self.client.patch(self.base+'/stages/2',json={'revision':1,'data':stage['data'],'reason':'check recomputation'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['data']['candidates'][0]['length_mm'],8000)
        response=self.client.post(self.base+'/stages/2/approve',json={'revision':1})
        self.assertEqual(response.status_code,409)

    def test_pending_requires_explicit_acknowledgement(self):
        stage=self.stage(1)
        stage['data']['observations'][0]['end_boundary']['kind']='UNKNOWN'
        edited=self.client.patch(self.base+'/stages/1',json={'revision':1,'data':stage['data'],'reason':'endpoint unclear'})
        self.assertEqual(edited.status_code,200,edited.text)
        self.approve(1)
        response=self.client.post(self.base+'/stages/2/approve',json={'revision':1})
        self.assertEqual(response.status_code,422)
        self.assertEqual(self.stage(2)['status'],'awaiting_review')

    def test_upstream_edit_restarts_and_hides_previous_exports(self):
        for n in (1,2,3):self.approve(n)
        stage=self.stage(1)
        stage['data']['observations'][0]['deferred']=True
        response=self.client.patch(self.base+'/stages/1',json={'revision':1,'data':stage['data'],'reason':'review endpoint'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.stage(2)['status'],'stale')
        self.assertEqual(self.client.get(self.base+'/download').status_code,404)
        self.approve(1)
        self.assertEqual(self.stage(2)['status'],'awaiting_review')
        self.assertFalse(self.stage(2)['data']['candidates'][0]['confirmed'])

    def test_evidence_and_file_access(self):
        evidence=self.client.get(self.base+'/evidence',params={'page_id':'f000-p001','q':'8000'}).json()
        self.assertEqual(evidence['total'],1)
        self.assertEqual(self.client.get(self.base+'/views/f000-p001.jpg').status_code,200)
        self.assertEqual(self.client.get(self.base+'/views/nonexistent.jpg').status_code,404)
        stage=self.stage(1)
        stage['data']['views'][0]['id']='../escape'
        response=self.client.patch(self.base+'/stages/1',json={'revision':1,'data':stage['data'],'reason':'invalid view id'})
        self.assertEqual(response.status_code,422)


class NativePdfTest(unittest.TestCase):
    def test_model_extra_boundary_note_is_preserved_without_changing_coordinates(self):
        views=prepare_mock_data()['views']
        original_box=list(views[0]['bbox'])
        views[0]['structure_bbox_hint']='钢架立面范围'
        result=parse_model_content(json.dumps({'views':views}),ViewResponse,{'id':'p'})
        self.assertEqual(result['views'][0]['bbox'],original_box)
        self.assertIn('钢架立面范围',result['views'][0]['boundary_reason'])
        self.assertNotIn('structure_bbox_hint',result['views'][0])

    def test_truncated_flash_response_is_preserved_and_retried_once(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'drawing.pdf';path.write_bytes(pdf_bytes())
            client=MagicMock();client.base_url='https://example.test'
            truncated=ChatCompletion(id='cut',created=0,model='deepseek-flash',object='chat.completion',
                choices=[{'index':0,'finish_reason':'length','message':{'role':'assistant','content':'{"views":['}}])
            good=ChatCompletion(id='good',created=0,model='deepseek-flash',object='chat.completion',
                choices=[{'index':0,'finish_reason':'stop','message':{'role':'assistant','content':json.dumps({'views':prepare_mock_data()['views']})}}])
            client.chat.completions.create.side_effect=[truncated,good]
            output=Path(tmp)/'raw.json'
            result=model_call(client,'deepseek-flash','test',{'id':'p','source_path':str(path),'page':1},output,ViewResponse)
            self.assertTrue(result['views'])
            actual_prompt=client.chat.completions.create.call_args.kwargs['messages'][1]['content'][0]['text']
            self.assertIn('HEA180 /EPB',actual_prompt)
            self.assertIn('你在第一阶段定位各独立视图',actual_prompt)
            self.assertIn('takeoff_role',actual_prompt)
            self.assertEqual(actual_prompt.count('你在第一阶段定位各独立视图'),1)
            self.assertTrue(output.with_suffix('.truncated.json').exists())
            self.assertEqual(client.chat.completions.create.call_count,2)
            options=client.chat.completions.create.call_args.kwargs
            self.assertEqual(options['extra_body']['thinking']['type'],'enabled')
            self.assertEqual(options['reasoning_effort'],'max')
            self.assertEqual(options['max_tokens'],131072)

    def test_boundary_review_requires_explanation_and_existing_evidence(self):
        schema=model_schema(BoundaryResponse)
        self.assertEqual(schema['$defs']['Boundary']['required'],['kind','evidence_ids','note'])
        observations=prepare_mock_data()['observations']
        identity=observations[0]['id']
        observations[0]['start_boundary']={'kind':'UNKNOWN','evidence_ids':[],'note':''}
        payload=BoundaryResponse(reviews=[{'observation_id':identity,
            'start_boundary':{'kind':'END','evidence_ids':['missing'],'note':'end connection'},
            'end_boundary':{'kind':'UNKNOWN','evidence_ids':[],'note':'not visible'}}]).model_dump()
        apply_boundary_reviews(observations,payload,{}, {})
        self.assertEqual(observations[0]['start_boundary']['kind'],'UNKNOWN')
        payload['reviews'][0]['start_boundary']['evidence_ids']=['end-vector']
        apply_boundary_reviews(observations,payload,{'end-vector':{'kind':'VECTOR'}},{})
        self.assertEqual(observations[0]['start_boundary']['kind'],'END')
        payload['reviews'].append(payload['reviews'][0])
        with self.assertRaises(ValueError):apply_boundary_reviews(observations,payload,{}, {})

    def test_wire_schema_does_not_request_server_owned_ids(self):
        schema=model_schema(ViewResponse)
        self.assertNotIn('id',schema['$defs']['View']['properties'])
        schema=model_schema(AnalysisResponse)
        self.assertNotIn('view_id',schema['$defs']['Observation']['properties'])
        self.assertEqual(request_options('test'),{})
        self.assertEqual(request_options('deepseek-flash')['reasoning_effort'],'high')
        self.assertEqual(request_options('deepseek-flash')['max_tokens'],65536)
        self.assertEqual(request_options('deepseek-flash','max')['reasoning_effort'],'max')
        with self.assertRaises(ValueError):request_options('deepseek-flash','low')

    def test_max_truncation_fails_without_disabling_thinking(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'drawing.pdf';path.write_bytes(pdf_bytes())
            client=MagicMock();client.base_url='https://example.test'
            truncated=ChatCompletion(id='cut',created=0,model='deepseek-flash',object='chat.completion',
                choices=[{'index':0,'finish_reason':'length','message':{'role':'assistant','content':'{"views":['}}])
            client.chat.completions.create.return_value=truncated
            output=Path(tmp)/'raw.json'
            with self.assertRaisesRegex(ValueError,'截断'):
                model_call(client,'deepseek-flash','test',{'id':'p','source_path':str(path),'page':1},output,ViewResponse)
            self.assertEqual(client.chat.completions.create.call_count,2)
            self.assertTrue(output.with_suffix('.truncated.json').exists())
            self.assertTrue(output.with_suffix('.max.truncated.json').exists())
            efforts=[call.kwargs['reasoning_effort'] for call in client.chat.completions.create.call_args_list]
            self.assertEqual(efforts,['high','max'])
            self.assertTrue(all(call.kwargs['extra_body']['thinking']['type']=='enabled'
                                for call in client.chat.completions.create.call_args_list))

    def test_only_unique_vector_aliases_are_repaired(self):
        evidence={
            'p:v1.0.0':{'kind':'VECTOR'},
            'p:v2.0.0':{'kind':'VECTOR'},
            'p:v2.1.0':{'kind':'VECTOR'},
            'p:t1':{'kind':'TEXT'},
        }
        payload={'evidence_ids':['p:v1','p:v2','p:missing','p:t1']}
        missing=repair_model_evidence(payload,evidence,evidence_aliases(evidence))
        self.assertEqual(payload['evidence_ids'],['p:v1.0.0','p:t1'])
        self.assertEqual(missing,['p:v2','p:missing'])

    def test_model_view_ids_are_assigned_before_strict_validation(self):
        views=prepare_mock_data()['views']
        views[0]['id']='f001-p001-plan-107.5'
        views[1]['id']='../model-supplied/name'
        views[0]['page_id']='untrusted-page'
        result=parse_model_content(json.dumps({'views':views}),ViewResponse,{'id':'f001-p001'})
        self.assertEqual([v['id'] for v in result['views']],['f001-p001-view01','f001-p001-view02'])
        self.assertTrue(all(v['page_id']=='f001-p001' for v in result['views']))
        views[0]['bbox']=[.8,0,.2,1]
        with self.assertRaises(ValueError):
            parse_model_content(json.dumps({'views':views}),ViewResponse,{'id':'f001-p001'})

    def test_model_observation_ids_are_server_owned(self):
        data=prepare_mock_data()
        observations=data['observations']
        observations[0]['id']='model.id'
        observations[0]['view_id']='wrong-view'
        result=parse_model_content(json.dumps({'frame':data['frame'],'observations':observations}),
                                   AnalysisResponse,{'id':'f001-p001-view01'})
        self.assertEqual(result['observations'][0]['id'],'f001-p001-view01-o001')
        self.assertTrue(all(o['view_id']=='f001-p001-view01' for o in result['observations']))

    def test_redundant_fixed_direction_is_removed_without_dropping_level_planes(self):
        views=prepare_mock_data()['views']
        views[0].update(visible_axes=['X','Y','Z'],fixed_axis='Z',
            axis_marks=[{'label':'EL.107.500(TOS)','direction':'Z','point':[.2,.2],'evidence_ids':['t']}])
        views[1].update(visible_axes=['Z','Y','X'],fixed_axis='Y')
        audit=MagicMock()
        result=parse_model_content(json.dumps({'views':views}),ViewResponse,{'id':'p'},audit)
        self.assertEqual(result['views'][0]['visible_axes'],['X','Y'])
        self.assertEqual(result['views'][1]['visible_axes'],['Z','X'])
        self.assertEqual(result['views'][0]['axis_marks'],views[0]['axis_marks'])
        self.assertEqual(result['views'][0]['bbox'],views[0]['bbox'])
        self.assertTrue(any(call.args[0]=='VIEW_VISIBLE_AXES_NORMALIZED' for call in audit.event.call_args_list))
        # Without an explicit fixed direction the parser must not invent one.
        views[0]['fixed_axis']=None
        with self.assertRaises(ValueError):
            parse_model_content(json.dumps({'views':views}),ViewResponse,{'id':'p'})
        views[0].update(visible_axes=['X','Z'],fixed_axis='Z')
        with self.assertRaises(ValueError):
            parse_model_content(json.dumps({'views':views}),ViewResponse,{'id':'p'})

    def test_resume_revalidates_saved_raw_response_without_requesting_model(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory=Path(tmp)
            upload=directory/'drawing.pdf';upload.write_bytes(pdf_bytes())
            saved=directory/'responses'/'attempt-old'/'f001-p001_views.json'
            saved.parent.mkdir(parents=True)
            views=prepare_mock_data()['views']
            views[0]['id']='f001-p001-plan-107.5'
            raw=ChatCompletion(id='saved',created=0,model='test',object='chat.completion',
                choices=[{'index':0,'finish_reason':'stop','message':{'role':'assistant','content':json.dumps({'views':views})}}])
            saved.write_text(raw.model_dump_json(),encoding='utf-8')
            target=directory/'responses'/'attempt-new'/'f001-p001_views.json'
            client=MagicMock();client.base_url='https://example.test'
            events=[]
            entry={'id':'f001-p001','source_path':str(upload),'page':1}
            result=model_call(client,'test','same prompt',entry,target,ViewResponse,reuse_saved=True,emit=events.append)
            client.chat.completions.create.assert_not_called()
            self.assertEqual(result['views'][0]['id'],'f001-p001-view01')
            self.assertEqual(saved.read_bytes(),target.read_bytes())
            self.assertTrue(events)
            # A signed response for another request must not be reused.
            saved.unlink()
            client.chat.completions.create.side_effect=RuntimeError('new request required')
            with self.assertRaisesRegex(RuntimeError,'new request required'):
                model_call(client,'test','changed prompt',entry,directory/'responses'/'attempt-third'/target.name,
                           ViewResponse,reuse_saved=True)

    def test_native_titles_ignore_reference_tables_and_keep_vectors(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'drawing.pdf';path.write_bytes(pdf_bytes())
            page=extract_page({'id':'p','source_path':str(path),'page':1,'file':path.name})
            self.assertEqual([t['kind'] for t in page['titles']],['PLAN','ELEVATION'])
            self.assertTrue(page['vectors'])
            self.assertTrue(page['vectors'][0]['points'])

    def test_max_resume_cannot_reuse_unsigned_deepseek_response(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);upload=root/'drawing.pdf';upload.write_bytes(pdf_bytes())
            saved=root/'responses/attempt-old/page_views.json';saved.parent.mkdir(parents=True)
            raw=ChatCompletion(id='saved',created=0,model='deepseek-flash',object='chat.completion',
                choices=[{'index':0,'finish_reason':'stop','message':{'role':'assistant',
                          'content':json.dumps({'views':prepare_mock_data()['views']})}}])
            saved.write_text(raw.model_dump_json(),encoding='utf-8')
            client=MagicMock();client.base_url='https://example.test'
            client.chat.completions.create.side_effect=RuntimeError('new max request required')
            with self.assertRaisesRegex(RuntimeError,'new max request required'):
                model_call(client,'deepseek-flash','same prompt',{'id':'page','source_path':str(upload),'page':1},
                           root/'responses/attempt-new/page_views.json',ViewResponse,reuse_saved=True,reasoning_effort='max')
            self.assertEqual(client.chat.completions.create.call_args.kwargs['reasoning_effort'],'max')

    def test_scan_like_pdf_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'blank.pdf'
            doc=fitz.open();doc.new_page();doc.save(path);doc.close()
            with self.assertRaises(ValueError):extract_page({'id':'p','source_path':str(path),'page':1,'file':path.name})

    def test_model_raw_response_survives_invalid_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'drawing.pdf';path.write_bytes(pdf_bytes())
            client=MagicMock()
            client.chat.completions.create.return_value=ChatCompletion(id='test',created=0,model='test',object='chat.completion',
                choices=[{'index':0,'finish_reason':'stop','message':{'role':'assistant','content':'not-json'}}])
            response_path=Path(tmp)/'raw.json'
            with self.assertRaises(ValueError):
                model_call(client,'test','test',{'source_path':str(path),'page':1},response_path,ViewResponse)
            self.assertIn('not-json',response_path.read_text())

    def test_five_real_sample_pdf_view_coverage(self):
        folder=Path(os.getenv('LLM_SAMPLE_DIR',str(Path.home()/'Desktop'/'TeklaProgram'/'建模图纸资料'/'建模图纸资料')))
        files=[folder/f'4422-AI-DU-302003PR01SA-0{n}_00.pdf' for n in range(1,6)]
        if not all(p.exists() for p in files):self.skipTest('设置 LLM_SAMPLE_DIR 后运行真实图纸检查。')
        titles=[];text=[];native_pages=[]
        for i,path in enumerate(files):
            page=extract_page({'id':f'f{i:03d}-p001','source_path':str(path),'page':1,'file':path.name})
            native_pages.append(page)
            titles.extend(page['titles']);text.extend(e['text'] for e in page['texts'])
        self.assertEqual(sum(t['kind']=='PLAN' for t in titles),5)
        self.assertEqual(sum(t['kind']=='ELEVATION' for t in titles),7)
        self.assertEqual(sum(t['kind']=='ISOMETRIC' for t in titles),1)
        self.assertEqual({t['elevation_m'] for t in titles if t['kind']=='PLAN'}, {100,107.5,110.5,115,119})
        self.assertTrue(any('EL.103.800' in t for t in text))
        self.assertTrue(any('EL.111.500' in t for t in text))
        self.assertTrue(any('-100mm BELOW TOS' in t for t in text))
        self.assertIn('480',text)
        self.assertIn('925',text)
        inventory=build_inventory({'views':[],'observations':[]},native_pages)
        self.assertEqual(inventory['profile_count'],16)
        self.assertEqual(inventory['view_counts'],{'PLAN':5,'ELEVATION':7,'ISOMETRIC':1})
        self.assertEqual({p['profile'] for p in inventory['profiles']},
            {'HEA160','HEA180','HEA200','HEA240','HEA260','HEA300','HEA340','HEA360',
             'HEA450','HEA500','HEB550','IPET270','IPET300','L80*8','L100*10','L120*12'})
