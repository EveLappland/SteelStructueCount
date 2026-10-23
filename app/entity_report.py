"""Entity ledger, observation audit and single-component expansion derived once."""
from collections import Counter
import csv
import html
import json

from openpyxl import Workbook


def component_rows(entities):
    rows=[]
    for e in entities:
        for member in range(e['member_cnt']):
            for component in range(e['component_multiplier']):
                rows.append({'component_id':f"{e['id']}-{member+1}-{component+1}",'entity_id':e['id'],
                    'profile':e['profile'],'length_mm':e['length_mm'],'member_cnt':1,
                    'view_tags':e['view_tags'],'status':e['status']})
    return rows


def export(directory,report,observations):
    entities=report['entities']; components=component_rows(entities)
    counts=Counter((c['profile'],c['length_mm']) for c in components)
    summary=[{'profile':p,'length_mm':length,'quantity':n,
              'total_length_mm':round(length*n,2) if length is not None else None}
             for (p,length),n in sorted(counts.items(),key=lambda item:(item[0][0],item[0][1] is None,item[0][1] or 0))]
    def cell(value):
        return json.dumps(value,ensure_ascii=False) if isinstance(value,(list,dict)) else value
    def table(sheet,rows,fields):
        sheet.append(fields)
        for row in rows: sheet.append([cell(row.get(f)) for f in fields])
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    entity_fields=['id','spec','length_mm','member_cnt','component_multiplier','view_tags','method','status',
                   'observation_ids','length_source_id','reason','measurement']
    with (directory/'entities.csv').open('w',encoding='utf-8-sig',newline='') as target:
        writer=csv.writer(target); writer.writerow(entity_fields)
        for row in entities: writer.writerow([cell(row.get(f)) for f in entity_fields])
    book=Workbook(); book.active.title='实体台账'
    table(book.active,entities,entity_fields)
    table(book.create_sheet('单根规格长度汇总'),summary,['profile','length_mm','quantity','total_length_mm'])
    compatible_summary=book.create_sheet('Quantity Length Summary')
    compatible_summary.append(['规格','名义长度 mm','单根数量','总长度 mm'])
    for row in summary:
        compatible_summary.append([row['profile'],row['length_mm'],row['quantity'],row['total_length_mm']])
    compatible_summary.freeze_panes='A2'; compatible_summary.auto_filter.ref=compatible_summary.dimensions
    table(book.create_sheet('展开单根台账'),components,['component_id','entity_id','profile','length_mm','view_tags','status'])
    table(book.create_sheet('跨视图观察对应'),report['observation_entity_map'],[
        'observation_id','label_id','view_uid','entity_id','original_length_mm','entity_length_mm','changed'])
    table(book.create_sheet('原始观察'),list(observations.values()),['observation_id','view_title','label_id','spec',
        'view_tags','length_mm','endpoints','basis','issue','status','selected_role'])
    table(book.create_sheet('自动任务审计'),report['cases'],['id','kind','state','error','observation_ids'])
    table(book.create_sheet('未知与回退'),[e for e in entities if e['length_mm'] is None or e['method']=='FALLBACK'],entity_fields)
    book.create_sheet('范围').append(['实体/组合先去重，单根台账仅在导出时展开一次。UNKNOWN保留数量；未完成匹配回退为独立观察，可能仍有重复。数量守恒和算术通过不代表工程准确率。'])
    book.save(directory/'stage_three.xlsx')
    write=lambda name,value:(directory/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    write('components.json',components); write('quantity_summary.json',summary)
    esc=lambda x:html.escape(str(x),quote=True)
    rows=[]
    for e in entities:
        sources=[]
        for identity in e['observation_ids']:
            r=observations[identity]
            sources.append(f"{r['view_title']} / {r['label_id']} / 原长度 {r['length_mm']}")
        rows.append('<tr>'+''.join(f'<td>{esc(x)}</td>' for x in (e['spec'],e['length_mm'] if e['length_mm'] is not None else 'UNKNOWN',
            e['member_cnt'],e['view_tags'],e['method']))+
            f'<td><details><summary>{len(sources)}条来源观察 · {esc(e["status"])}</summary><pre>{esc(chr(10).join(sources))}</pre>'
            f'<p>{esc(e["reason"])}</p><pre>{esc(json.dumps(e["measurement"],ensure_ascii=False,indent=2))}</pre></details></td></tr>')
    pictures=''.join(f'<details><summary>{esc(v["title"])}</summary><a href="{esc(v["image_file"])}">打开高清原图</a>'
        f'<img loading="lazy" src="{esc(v["image_file"])}" style="max-width:100%"></details>' for v in report['views'])
    tasks=''.join(f'<tr><td>{esc(c["kind"])}</td><td>{esc(c["id"])}</td><td>{esc(c["state"])}</td>'
        f'<td>{esc(c.get("error") or "")}</td></tr>' for c in report['cases'])
    page=f'''<!doctype html><html lang="zh"><meta charset="utf-8"><title>第三阶段 · 跨视图实体归约</title>
<style>body{{font-family:system-ui;margin:24px;color:#183047;line-height:1.6}}table{{border-collapse:collapse;width:100%;font-size:14px}}
th,td{{border:1px solid #ccd6df;padding:8px;vertical-align:top;text-align:left}}th{{background:#eef4f9}}pre{{white-space:pre-wrap;max-width:700px}}</style>
<h1>第三阶段 · 跨视图关联、长度仲裁与实体归约</h1>
<p>保留全部原始标签；同一视图不同标签不压为一根。匹配组数量保留，组内无法区分时不编造逐根配对。自动完成，无人工修改台账。</p>
<p>{esc(report['summary'])}</p><p>{esc(report['conservation'])}</p>
<p><a href="stage_three.xlsx">下载Excel</a> · <a href="entities.csv">实体CSV</a> · <a href="stage_three.json">完整JSON</a> · <a href="quantity_summary.json">展开单根汇总</a></p>
<p>UNKNOWN保留数量。FALLBACK表示未完成自动裁决，保留独立观察，可能仍有重复；已填数值也不等于工程准确率。</p>
<h2>来源原图</h2>{pictures}<h2>实体/组合台账（尚未展开组合）</h2>
<table><tr><th>规格</th><th>名义长度mm</th><th>实体数量</th><th>所属面</th><th>处理方式</th><th>来源与依据</th></tr>{''.join(rows)}</table>
<h2>自动处理任务</h2><table><tr><th>类型</th><th>任务ID</th><th>状态</th><th>说明</th></tr>{tasks}</table></html>'''
    (directory/'index.html').write_text(page,encoding='utf-8')
