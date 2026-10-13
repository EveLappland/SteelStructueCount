from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = '143A5A'
TEAL = '087E8B'
SKY = 'DCECF5'
GREEN = 'DCEAD7'
AMBER = 'FCEBC5'


def _title(ws, text, columns):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = ws.cell(1, 1, text)
    cell.fill = PatternFill('solid', fgColor=NAVY)
    cell.font = Font(color='FFFFFF', bold=True, size=16)
    cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 28


def _header(ws, row, values):
    for index, value in enumerate(values, 1):
        cell = ws.cell(row, index, value)
        cell.fill = PatternFill('solid', fgColor=TEAL)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.auto_filter.ref = f'A{row}:{get_column_letter(len(values))}{row}'


def _widths(ws, widths):
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = 'A4'


def write_workbook(directory: Path, raw_records: list[dict], records: list[dict], include_lengths: bool, summary: str,
                   review_queue: list[dict], review_decisions: list[dict], change_log: list[dict]):
    output = directory / 'output'
    output.mkdir(exist_ok=True)
    path = output / 'steel_takeoff.xlsx'
    workbook = Workbook()
    summary_ws = workbook.active
    summary_ws.title = 'Reconciled Summary'
    _title(summary_ws, 'Reconciled Structural Steel Quantity and Length Summary', 6)
    summary_ws.merge_cells('A2:F2')
    summary_ws['A2'] = 'PDF-derived result. Model interpretation requires engineering review before procurement.'
    summary_ws['A2'].font = Font(italic=True, color='617386')
    _header(summary_ws, 4, ['PROFILE', 'MATERIAL', 'QUANTITY (pcs)', 'COUNT STATUS', 'TOTAL LENGTH (mm)', 'NOTES'])

    grouped = defaultdict(list)
    confirmed = [record for record in records if record['status'] == 'CONFIRMED' and record['decision'] != 'UNRESOLVED' and record['profile'] and isinstance(record['quantity'], int) and not isinstance(record['quantity'], bool) and record['quantity'] > 0]
    for record in confirmed:
        grouped[(record['profile'], record['material'])].append(record)
    total = 0
    for row, ((profile, material), items) in enumerate(sorted(grouped.items()), 5):
        quantity = sum(item['quantity'] for item in items)
        total += quantity
        unresolved = any(item['status'] != 'CONFIRMED' for item in items)
        missing_lengths = include_lengths and any(not item.get('length_mm') for item in items)
        length_total = '' if missing_lengths else sum(item['length_mm'] * item['quantity'] for item in items)
        notes = '; '.join(dict.fromkeys(item['role'] for item in items if item['role']))[:500]
        if missing_lengths:
            notes = '; '.join(filter(None, [notes, '部分构件未取得可核验长度，未将其按零计入总长度。']))
        values = [profile, material, quantity, 'PARTIAL' if unresolved else 'CONFIRMED', length_total if include_lengths else '', notes]
        for column, value in enumerate(values, 1):
            cell = summary_ws.cell(row, column, value)
            if row % 2:
                cell.fill = PatternFill('solid', fgColor=SKY)
            if column == 4:
                cell.fill = PatternFill('solid', fgColor=AMBER if unresolved else GREEN)
                cell.font = Font(bold=True)
    total_row = 5 + len(grouped)
    summary_ws.cell(total_row, 1, 'CONFIRMED TOTAL')
    summary_ws.cell(total_row, 3, total)
    for column in range(1, 7):
        cell = summary_ws.cell(total_row, column)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
    _widths(summary_ws, [18, 16, 17, 18, 20, 58])

    audit_ws = workbook.create_sheet('Count Audit')
    _title(audit_ws, 'Raw Physical Member Count and Length Audit', 16)
    _header(audit_ws, 4, ['MEMBER ID', 'PROFILE', 'MATERIAL', 'QUANTITY', 'DRAWN BRACE LINES', 'ANGLE MULTIPLIER', 'NOMINAL LENGTH (mm)', 'TOTAL LENGTH (mm)', 'LENGTH METHOD', 'SOURCE', 'LEVEL', 'LOCATION', 'PHYSICAL SPANS', 'SPLICE EVIDENCE', 'ROLE', 'PASS 2 DECISION / EVIDENCE'])
    for row, record in enumerate(raw_records, 5):
        length_mm = record.get('length_mm') or ''
        total_length = length_mm * record['quantity'] if include_lengths and length_mm and isinstance(record.get('quantity'), int) else ''
        physical_spans = '\n'.join(str(item) for item in record.get('physical_spans', []))
        values = [record['member_id'], record['profile'], record['material'], record['quantity'], record.get('drawn_brace_lines', ''), record.get('quantity_multiplier', 1), length_mm, total_length, record.get('length_method', 'UNRESOLVED'), record['source'], record['level'], record['location'], physical_spans, record.get('splice_evidence', ''), record['role'], f"{record['decision']}; {record['notes']}"]
        for column, value in enumerate(values, 1):
            audit_ws.cell(row, column, value).alignment = Alignment(vertical='top', wrap_text=True)
    _widths(audit_ws, [20, 16, 14, 12, 18, 17, 20, 20, 20, 24, 20, 32, 42, 28, 25, 52])

    queue_ws = workbook.create_sheet('Review Queue')
    _title(queue_ws, 'Records Selected for Targeted Reconciliation', 7)
    _header(queue_ws, 4, ['PROFILE', 'MEMBER ID', 'CURRENT QUANTITY', 'CURRENT LENGTH (mm)', 'SOURCE', 'ROLE', 'SCREENING REASONS'])
    row = 5
    for group in review_queue:
        reason = ' '.join(group.get('reasons', []))
        for record in group.get('records', []):
            values = [group.get('profile', ''), record.get('member_id', ''), record.get('quantity', ''), record.get('length_mm', ''), record.get('source', ''), record.get('role', ''), reason]
            for column, value in enumerate(values, 1):
                queue_ws.cell(row, column, value).alignment = Alignment(vertical='top', wrap_text=True)
            row += 1
    _widths(queue_ws, [16, 20, 18, 20, 34, 28, 68])

    decisions_ws = workbook.create_sheet('Review Decisions')
    _title(decisions_ws, 'Targeted Model Review Decisions', 10)
    _header(decisions_ws, 4, ['MEMBER ID', 'ACTION', 'REASON / VIEW EVIDENCE', 'PROFILE CORRECTION', 'ROLE CORRECTION', 'QUANTITY CORRECTION', 'RETAINED PHYSICAL SPANS', 'SPLICE EVIDENCE', 'LENGTH CORRECTION (mm)', 'LENGTH METHOD'])
    for row, decision in enumerate(review_decisions, 5):
        physical_spans = decision.get('physical_spans', [])
        values = [decision.get('member_id', ''), decision.get('action', ''), decision.get('reason', ''), decision.get('profile', ''), decision.get('role', ''), decision.get('quantity', ''), '\n'.join(str(item) for item in physical_spans) if isinstance(physical_spans, list) else '', decision.get('splice_evidence', ''), decision.get('length_mm', ''), decision.get('length_method', '')]
        for column, value in enumerate(values, 1):
            decisions_ws.cell(row, column, value).alignment = Alignment(vertical='top', wrap_text=True)
    _widths(decisions_ws, [20, 22, 68, 22, 28, 22, 42, 28, 24, 24])

    changes_ws = workbook.create_sheet('Change Log')
    _title(changes_ws, 'Reconciliation Change Log', 7)
    _header(changes_ws, 4, ['MEMBER ID', 'ACTION', 'QUANTITY BEFORE', 'QUANTITY AFTER', 'LENGTH BEFORE (mm)', 'LENGTH AFTER (mm)', 'REASON'])
    for row, change in enumerate(change_log, 5):
        values = [change.get('member_id', ''), change.get('action', ''), change.get('before_quantity', ''), change.get('after_quantity', ''), change.get('before_length_mm', ''), change.get('after_length_mm', ''), change.get('reason', '')]
        for column, value in enumerate(values, 1):
            changes_ws.cell(row, column, value).alignment = Alignment(vertical='top', wrap_text=True)
    _widths(changes_ws, [20, 22, 20, 20, 22, 22, 70])

    exceptions = [item for item in records if item not in confirmed]
    if exceptions:
        exceptions_ws = workbook.create_sheet('Exceptions')
        _title(exceptions_ws, 'Items Requiring Engineering Confirmation', 5)
        _header(exceptions_ws, 4, ['MEMBER ID', 'PROFILE / MATERIAL', 'SOURCE', 'QUESTION', 'NOTES'])
        for row, record in enumerate(exceptions, 5):
            values = [record['member_id'], f"{record['profile']} / {record['material']}", record['source'], record['notes'] or 'Model evidence is insufficient.', record['location']]
            for column, value in enumerate(values, 1):
                exceptions_ws.cell(row, column, value).alignment = Alignment(vertical='top', wrap_text=True)
        _widths(exceptions_ws, [22, 24, 28, 58, 35])

    if include_lengths:
        length_audit = workbook.create_sheet('Length Audit')
        _title(length_audit, 'Locked Counts and Independent Length Verification', 7)
        _header(length_audit, 4, ['MEMBER ID', 'PROFILE', 'QUANTITY', 'COUNT STATUS', 'LENGTH (mm)', 'LENGTH STATUS', 'EVIDENCE / NOTES'])
        for row, record in enumerate(records, 5):
            values = [record['member_id'], record['profile'], record.get('quantity'), record['status'],
                      record.get('length_mm'), record.get('length_status', 'UNRESOLVED'), record.get('notes', '')]
            for column, value in enumerate(values, 1):
                length_audit.cell(row, column, value).alignment = Alignment(vertical='top', wrap_text=True)
        _widths(length_audit, [28, 18, 14, 20, 20, 20, 70])
        schedule_ws = workbook.create_sheet('Profile Length Schedule')
        _title(schedule_ws, 'Profile Length Schedule - One Row Per Physical Piece', 8)
        _header(schedule_ws, 4, ['PROFILE', 'MATERIAL', 'LENGTH (mm)', 'PIECE ID', 'SOURCE', 'LEVEL / INTERVAL', 'LENGTH METHOD', 'NOTES'])
        row = 5
        for record in confirmed:
            for piece_number in range(1, record['quantity'] + 1):
                values = [record['profile'], record['material'], record.get('length_mm') or '', f"{record['member_id']}-{piece_number:03d}", record['source'], record['level'], record.get('length_method', 'UNRESOLVED'), record['notes']]
                for column, value in enumerate(values, 1):
                    schedule_ws.cell(row, column, value).alignment = Alignment(vertical='top', wrap_text=True)
                row += 1
        _widths(schedule_ws, [16, 14, 16, 24, 26, 23, 22, 52])

    notes_ws = workbook.create_sheet('Method Notes')
    _title(notes_ws, 'Method Notes', 2)
    notes_ws['A3'] = summary
    notes_ws['A5'] = 'Only confirmed records with valid profile and positive integer quantity contribute to totals and piece length schedules. Unresolved records remain in Count Audit and Exceptions; unknown quantities are blank, never defaulted to 1.'
    notes_ws['A5'].alignment = Alignment(wrap_text=True, vertical='top')
    notes_ws['A3'].alignment = Alignment(wrap_text=True, vertical='top')
    notes_ws.column_dimensions['A'].width = 115
    workbook.save(path)
    return path, total, len(exceptions)
